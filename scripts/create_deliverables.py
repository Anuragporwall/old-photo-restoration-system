"""Create editable DOCX report and PPTX presentation from the project content."""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
DOCS, ASSETS = ROOT / "docs", ROOT / "assets"

NAVY, BLUE, CYAN, WHITE, MUTED = "0B1020", "1A3159", "25B9D7", "F4F8FF", "9FB0CA"


def add_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
    title = doc.add_heading("Old Photograph Restoration System", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Using Digital Image Processing\nProject Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs: run.font.size = Pt(14); run.font.color.rgb = RGBColor(35, 92, 145)
    doc.add_paragraph("A classical Digital Image Processing project built with Python, Streamlit, OpenCV, and NumPy.")
    sections = [
        ("Abstract", "Old photographs often suffer from fading, low contrast, dust, impulse noise, and soft edges. This project provides a browser-based restoration system using only classical digital image processing. The system applies filtering, local contrast enhancement, sharpening, interpolation, and frequency-domain processing while preserving the original image and supporting comparison of named variants."),
        ("Objectives", "Restore moderate image degradation without AI or machine learning. Preserve colour by default, provide both automatic and step-by-step workflows, keep per-photo variants, and prevent avoidable errors through validation and safe retry."),
        ("Algorithms", "Median filtering removes impulse noise. Gaussian smoothing suppresses fine random noise. CLAHE enhances local contrast on the CIELAB luminance channel. Unsharp masking and Laplacian enhancement recover edge detail. Bilinear interpolation resizes images. Low-pass and high-pass FFT masks demonstrate frequency-domain restoration."),
        ("System Design", "The Streamlit dashboard validates uploads, stores every photo and variant in temporary session state, invokes the OpenCV/NumPy restoration engine, then provides comparison, analysis, and PNG/ZIP export. The original image is retained separately from all results."),
        ("Safety and Reliability", "The app accepts at most five images and caps upload size at 20 MB. Dimension and output-pixel limits protect memory. Destructive actions require confirmation. If an automatic operation fails, the engine retries once with lower-impact settings and reports only the final outcome. Previous variants are never overwritten by a failure."),
        ("Evaluation", "Automated tests check shape preservation, filtering, Fourier filtering, metrics, retry behavior, and enlargement protection. Manual acceptance checks cover upload validation, session variants, confirmations, downloads, and continued batch use after an individual failure."),
        ("Conclusion", "The project demonstrates explainable, colour-aware photograph restoration with classical DIP. It gives students direct control over processing parameters while offering a safe balanced starting point for practical restoration."),
    ]
    for heading, body in sections:
        doc.add_heading(heading, 1); doc.add_paragraph(body)
    image = ASSETS / "sample_comparison.png"
    if image.exists():
        doc.add_heading("Sample Result", 1)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Inches(6.35))
        doc.add_paragraph("Figure 1. Synthetic faded input (left) and balanced restoration output (right).", style="Caption")
    doc.add_heading("Run Instructions", 1)
    doc.add_paragraph("Install the packages in requirements.txt, then run: streamlit run app.py")
    doc.save(DOCS / "project_report.docx")


def color(hexcode): return PRGBColor.from_string(hexcode)

def textbox(slide, text, x, y, w, h, size=22, colour=WHITE, bold=False):
    shape = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    paragraph = tf.paragraphs[0]; paragraph.text = text; paragraph.font.size = PPt(size); paragraph.font.bold = bold; paragraph.font.name = "Aptos"; paragraph.font.color.rgb = color(colour)
    return shape


def base_slide(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color(NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PInches(.14), PInches(7.5)); accent.fill.solid(); accent.fill.fore_color.rgb = color(CYAN); accent.line.fill.background()
    textbox(slide, title, .55, .35, 11.7, .55, 29, WHITE, True)
    textbox(slide, f"OLD PHOTO RESTORATION SYSTEM  |  {number:02d}", .58, 7.08, 5.6, .22, 8, MUTED)
    return slide


def bullet_slide(prs, title, bullets, number):
    slide = base_slide(prs, title, number)
    y = 1.25
    for text in bullets:
        block = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PInches(.65), PInches(y), PInches(11.35), PInches(.78))
        block.fill.solid(); block.fill.fore_color.rgb = color("111B30"); block.line.color.rgb = color("263B63")
        textbox(slide, text, .95, y + .16, 10.75, .42, 17, WHITE)
        y += .96
    return slide


def add_pptx():
    prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color(NAVY)
    textbox(slide, "Old Photograph\nRestoration System", .8, 1.1, 7.2, 1.7, 38, WHITE, True)
    textbox(slide, "Using Classical Digital Image Processing", .85, 3.02, 6.5, .45, 20, CYAN)
    textbox(slide, "Python  •  Streamlit  •  OpenCV  •  NumPy\nNo AI or machine learning", .85, 4.0, 5.0, .6, 16, MUTED)
    image = ASSETS / "sample_comparison.png"
    if image.exists(): slide.shapes.add_picture(str(image), PInches(7.1), PInches(1.15), width=PInches(5.3))
    bullet_slide(prs, "Problem and goals", ["Restore fading, noise, reduced contrast, and soft detail in old photographs.", "Use explainable classical DIP methods only.", "Preserve original images and let users compare named variants.", "Make experimentation safe for a multi-photo browser session."], 2)
    bullet_slide(prs, "Restoration pipeline", ["Median filter: removes impulse noise while keeping edges.", "Gaussian filter: reduces fine random noise.", "CLAHE on luminance: restores local contrast while preserving colour.", "Unsharp mask + Laplacian: improves visible edge detail."], 3)
    bullet_slide(prs, "Frequency-domain tools", ["2-D FFT shifts low frequencies to the spectrum centre.", "Low-pass mask suppresses fine noise; high-pass mask emphasizes detail.", "The dashboard visualizes the log-magnitude frequency spectrum.", "Per-channel filtering maintains colour image support."], 4)
    bullet_slide(prs, "User experience", ["Dark dashboard with up to five photos per session.", "Balanced automatic preset plus detailed step-by-step controls.", "Original/result side-by-side view, image metrics, and histogram.", "Editable PNG naming and ZIP download for every session variant."], 5)
    bullet_slide(prs, "Safety and reliability", ["Strict file type, upload size, decoded-pixel, and output-pixel validation.", "Confirmation gates for large output, variant deletion, photo deletion, and session clearing.", "Automatic retry uses safer settings and shows only the final outcome.", "A failure never replaces an existing result or blocks the other session photos."], 6)
    bullet_slide(prs, "Testing and conclusion", ["Automated tests validate filters, metrics, safe retry, and output limits.", "Manual checks validate upload, variants, comparisons, exports, and failure isolation.", "The result is an explainable teaching and restoration tool based on classical DIP.", "Future work: morphology-based scratch detection and illumination correction."], 7)
    prs.save(DOCS / "project_presentation.pptx")


if __name__ == "__main__":
    add_docx(); add_pptx()
